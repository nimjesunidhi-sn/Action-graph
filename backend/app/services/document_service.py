from pathlib import Path
import re, subprocess, tempfile
from app.core.config import settings

ALLOWED={'.txt','.md','.markdown','.pdf','.docx','.srt','.vtt','.mp3','.wav','.mp4','.m4a','.mov','.webm'}
def extract_text(path: Path) -> str:
    ext=path.suffix.lower()
    if ext in {'.txt','.md','.markdown','.srt','.vtt'}:
        text=path.read_text(encoding='utf-8-sig',errors='replace')
        if ext in {'.srt','.vtt'}:
            text=re.sub(r'(?m)^\s*\d+\s*$','',text); text=re.sub(r'(?m)^\s*(?:\d{2}:)?\d{2}:\d{2}[,.]\d{3}\s+-->.*$','',text); text=re.sub(r'(?m)^WEBVTT\s*$','',text)
        return text.strip()
    if ext=='.pdf':
        import fitz
        doc=fitz.open(path); text='\n'.join(page.get_text() for page in doc)
        if not text.strip(): raise ValueError('This PDF has no extractable text. Scanned PDFs need OCR before processing.')
        return text.strip()
    if ext=='.docx':
        from docx import Document
        doc=Document(path); return '\n'.join([p.text for p in doc.paragraphs]+[cell.text for table in doc.tables for row in table.rows for cell in row.cells]).strip()
    if ext in {'.mp3','.wav','.m4a','.mp4','.mov','.webm'}:
        try: from faster_whisper import WhisperModel
        except ImportError: raise ValueError('Audio/video transcription is optional. Install faster-whisper and FFmpeg, then retry this meeting.')
        audio_path=path; temp_path=None
        try:
            if ext in {'.mp4','.mov','.webm'}:
                temp=tempfile.NamedTemporaryFile(suffix='.wav',delete=False); temp_path=Path(temp.name); temp.close()
                try: subprocess.run(['ffmpeg','-y','-i',str(path),'-vn','-ac','1','-ar','16000',str(temp_path)],check=True,capture_output=True,timeout=600)
                except FileNotFoundError: raise ValueError('FFmpeg was not found. Install FFmpeg and make sure it is available on PATH.')
                except subprocess.CalledProcessError: raise ValueError('FFmpeg could not extract audio from this video. Check that the file contains a supported audio track.')
                audio_path=temp_path
            model=WhisperModel(settings.whisper_model,device='cpu',compute_type='int8')
            segments,_=model.transcribe(str(audio_path),vad_filter=True)
            return '\n'.join(f'[{int(s.start//3600):02d}:{int(s.start%3600//60):02d}:{s.start%60:06.3f}] {s.text.strip()}' for s in segments).strip()
        finally:
            if temp_path: temp_path.unlink(missing_ok=True)
    raise ValueError(f'{ext or "File"} is not supported for text extraction.')
